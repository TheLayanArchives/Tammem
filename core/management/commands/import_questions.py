
from collections import Counter
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from core.models import Question, QuestionOption


DATA_DIR = Path(__file__).resolve().parents[2] / "data"

MASTER_FILE = DATA_DIR / "TAMMEM_Master_Questions_By_Scope.xlsx"
TRANSLATION_FILE = DATA_DIR / "Questions_EN-AR (version 1).xlsx"
OPTIONS_FILE = DATA_DIR / "Tammem_Question_Option_Translations_Approved.xlsx"

TYPE_MAPPING = {
    "BOOLEAN": "BOOLEAN",
    "SINGLE_CHOICE": "SINGLE",
    "MULTI_SELECT": "MULTIPLE",
    "NUMBER": "NUMBER",
    "TEXT": "TEXT",
    "CURRENCY": "DECIMAL",
    "CURRENCY_RANGE": "SINGLE",

    # Temporary storage mappings for specialized inputs
    "COUNTRY_SELECT": "TEXT",
    "LOCATION_SELECT": "TEXT",
    "ACTIVITY_SELECT": "TEXT",
    "MULTI_ACTIVITY_SELECT": "TEXT",
    "PROFESSION_SELECT": "TEXT",
    "OWNER_PERCENTAGES": "TEXT",
}

REQUIRED_MAPPING = {
    "yes": "YES",
    "no": "NO",
    "conditional": "CONDITIONAL",
    "derived/conditional": "DERIVED_CONDITIONAL",
}

STATIC_TYPES = {
    "SINGLE_CHOICE",
    "MULTI_SELECT",
    "CURRENCY_RANGE",
}


def clean(value):
    return str(value).strip() if value is not None else ""


def read_sheet(path, sheet_name, required_headers):
    if not path.is_file():
        raise CommandError(f"File not found: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)

    try:
        if sheet_name not in wb.sheetnames:
            raise CommandError(
                f"Sheet '{sheet_name}' not found in {path.name}"
            )

        rows = iter(wb[sheet_name].values)
        headers = next(rows, None)

        if headers is None:
            raise CommandError(f"Empty sheet: {path.name}")

        headers = [clean(h) for h in headers]

        if len(headers) != len(set(headers)):
            raise CommandError(
                f"Duplicate column headers in {path.name}"
            )

        missing = set(required_headers) - set(headers)

        if missing:
            raise CommandError(
                f"Missing columns in {path.name}: {sorted(missing)}"
            )

        records = []

        for row_number, values in enumerate(rows, start=2):
            record = dict(zip(headers, values))
            record["_row_number"] = row_number
            records.append(record)

        return records

    finally:
        wb.close()


def index_questions(rows, source):
    indexed = {}

    for row in rows:
        code = clean(row.get("Question ID"))

        if not code:
            continue

        if code in indexed:
            raise CommandError(
                f"Duplicate question {code} in {source}"
            )

        indexed[code] = row

    return indexed


def static_options(input_type, source):
    if input_type not in STATIC_TYPES:
        return []

    source = clean(source)

    if not source:
        return []

    if source.lower().startswith("database-driven"):
        return []

    return [
        value.strip()
        for value in source.split(";")
        if value.strip()
    ]


class Command(BaseCommand):
    help = "Preview or import Tammem questions and approved option translations."

    def add_arguments(self, parser):
        parser.add_argument(
            "--apply",
            action="store_true",
            help="Write validated data to PostgreSQL.",
        )

    def handle(self, *args, **options):
        master_rows = read_sheet(
            MASTER_FILE,
            "Master Questions",
            [
                "Question ID",
                "Question Type",
                "Required",
                "Scope",
                "Options / Source",
            ],
        )

        translation_rows = read_sheet(
            TRANSLATION_FILE,
            "Sheet1",
            [
                "Question ID",
                "Question EN",
                "Question AR(Final Version)",
            ],
        )

        option_rows = read_sheet(
            OPTIONS_FILE,
            "Option Translations",
            [
                "Question ID",
                "Option Code",
                "Option EN",
                "Option AR",
                "Review Status",
                "Input Type",
            ],
        )

        master = index_questions(master_rows, "Master")
        translations = index_questions(
            translation_rows, "Translations"
        )

        if len(master) != 102:
            raise CommandError(
                f"Expected 102 questions, found {len(master)}"
            )

        if set(master) != set(translations):
            raise CommandError(
                "Question IDs do not match between master and translations. "
                f"Missing: {sorted(set(master) - set(translations))}; "
                f"Extra: {sorted(set(translations) - set(master))}"
            )

        approved_options = {}
        seen_option_keys = set()

        for row in option_rows:
            code = clean(row["Question ID"])
            option_code = clean(row["Option Code"])
            option_en = clean(row["Option EN"])
            option_ar = clean(row["Option AR"])
            status = clean(row["Review Status"])
            input_type = clean(row["Input Type"])

            key = (code, option_code)

            if key in seen_option_keys:
                raise CommandError(
                    f"Duplicate option: {code} / {option_code}"
                )

            seen_option_keys.add(key)

            if not code or not option_code:
                raise CommandError("Option with missing ID or code")

            if not option_en or not option_ar:
                raise CommandError(
                    f"Missing option translation: {key}"
                )

            if status != "Approved":
                raise CommandError(
                    f"Option not approved: {key}"
                )

            if code not in master:
                raise CommandError(
                    f"Option references unknown question: {code}"
                )

            if input_type != clean(
                master[code]["Question Type"]
            ):
                raise CommandError(
                    f"Input type mismatch: {key}"
                )

            approved_options[key] = {
                "label_en": option_en,
                "label_ar": option_ar,
            }

        prepared = []
        expected_keys = set()
        type_counts = Counter()

        for order, (code, row) in enumerate(
            master.items(), start=1
        ):
            input_type = clean(row["Question Type"])
            required = clean(row["Required"]).lower()
            scope = clean(row["Scope"])
            translation = translations[code]

            if input_type not in TYPE_MAPPING:
                raise CommandError(
                    f"{code}: unsupported type {input_type}"
                )

            if required not in REQUIRED_MAPPING:
                raise CommandError(
                    f"{code}: invalid required value {required}"
                )

            if scope not in {
                "GENERAL",
                "PROFILE",
                "ACTIVITY",
            }:
                raise CommandError(
                    f"{code}: invalid scope {scope}"
                )

            text_en = clean(translation["Question EN"])
            text_ar = clean(
                translation["Question AR(Final Version)"]
            )

            if not text_en or not text_ar:
                raise CommandError(
                    f"{code}: missing question translation"
                )

            labels = static_options(
                input_type,
                row["Options / Source"],
            )

            source = clean(row["Options / Source"])

            if (
                input_type in STATIC_TYPES
                and not labels
                and not source.lower().startswith("database-driven")
            ):
                raise CommandError(
                    f"{code}: missing static options"
                )

            question_options = []

            for index, label in enumerate(labels, start=1):
                option_code = f"EXCEL_{index:03d}"
                key = (code, option_code)
                expected_keys.add(key)

                translated = approved_options.get(key)

                if translated is None:
                    raise CommandError(
                        f"Missing approved option: {key}"
                    )

                if translated["label_en"] != label:
                    raise CommandError(
                        f"English option mismatch: {key}"
                    )

                question_options.append({
                    "code": option_code,
                    "label_en": label,
                    "label_ar": translated["label_ar"],
                    "value": label,
                    "display_order": index,
                    "is_active": True,
                })

            type_counts[input_type] += 1

            prepared.append({
                "code": code,
                "defaults": {
                    "text_en": text_en,
                    "text_ar": text_ar,
                    "question_type": TYPE_MAPPING[input_type],
                    "input_type": input_type,
                    "required_mode": REQUIRED_MAPPING[required],
                    "is_required": required == "yes",
                    "scope": scope,
                    "display_order": order,
                    "is_active": True,
                },
                "options": question_options,
            })

        if expected_keys != set(approved_options):
            extra = sorted(
                set(approved_options) - expected_keys
            )
            raise CommandError(
                f"Unexpected translated options: {extra}"
            )

        if len(expected_keys) != 109:
            raise CommandError(
                f"Expected 109 options, found {len(expected_keys)}"
            )

        existing_count = Question.objects.filter(
            code__in=list(master)
        ).count()

        self.stdout.write("\nVALIDATION SUCCESSFUL")
        self.stdout.write(
            f"Questions validated: {len(prepared)}"
        )
        self.stdout.write(
            f"Approved bilingual options: {len(expected_keys)}"
        )
        self.stdout.write(
            f"Existing matching questions: {existing_count}"
        )

        self.stdout.write("\nQuestion types:")
        for name, count in sorted(type_counts.items()):
            self.stdout.write(f"  {name}: {count}")

        if not options["apply"]:
            self.stdout.write(
                self.style.WARNING(
                    "\nPREVIEW ONLY: No database changes made."
                )
            )
            return

        created_questions = 0
        updated_questions = 0
        created_options = 0
        updated_options = 0

        with transaction.atomic():
            for item in prepared:
                question, created = Question.objects.update_or_create(
                    code=item["code"],
                    defaults=item["defaults"],
                )

                if created:
                    created_questions += 1
                else:
                    updated_questions += 1

                for option in item["options"]:
                    option_code = option["code"]

                    defaults = {
                        key: value
                        for key, value in option.items()
                        if key != "code"
                    }

                    _, option_created = (
                        QuestionOption.objects.update_or_create(
                            question=question,
                            code=option_code,
                            defaults=defaults,
                        )
                    )

                    if option_created:
                        created_options += 1
                    else:
                        updated_options += 1

        self.stdout.write(
            self.style.SUCCESS(
                "\nIMPORT COMPLETED SUCCESSFULLY"
            )
        )
        self.stdout.write(
            f"Questions created: {created_questions}"
        )
        self.stdout.write(
            f"Questions updated: {updated_questions}"
        )
        self.stdout.write(
            f"Options created: {created_options}"
        )
        self.stdout.write(
            f"Options updated: {updated_options}"
        )
