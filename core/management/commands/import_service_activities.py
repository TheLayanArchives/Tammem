
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from core.models import Activity, GovernmentService, ServiceActivity


CAFE_CODE_MAP = {
    "CAF": "CAFE_A",
    "TOB": "CAFE_B",
    "KSK": "CAFE_K",
    "CLD": "CAFE_C",
}


class Command(BaseCommand):
    help = "Import service-to-cafe activity mappings from Excel."

    def add_arguments(self, parser):
        parser.add_argument("excel_file", type=str)

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise CommandError(
                "openpyxl is missing. Install it using: pip install openpyxl"
            )

        file_path = Path(options["excel_file"])

        if not file_path.is_file():
            raise CommandError(f"Excel file not found: {file_path}")

        workbook = load_workbook(file_path, read_only=True, data_only=True)

        try:
            sheet = next(
                (
                    workbook[name]
                    for name in workbook.sheetnames
                    if {"Service_Code", "Cafe_Types"}.issubset(
                        {
                            str(cell.value).strip()
                            for cell in workbook[name][1]
                            if cell.value is not None
                        }
                    )
                ),
                None,
            )

            if sheet is None:
                raise CommandError(
                    "No sheet contains Service_Code and Cafe_Types columns."
                )

            rows = sheet.iter_rows(values_only=True)
            headers = [
                str(value).strip() if value is not None else ""
                for value in next(rows)
            ]

            code_index = headers.index("Service_Code")
            types_index = headers.index("Cafe_Types")

            mappings = set()
            service_codes = set()

            for row_number, row in enumerate(rows, start=2):
                raw_code = row[code_index]
                raw_types = row[types_index]

                if raw_code is None and raw_types is None:
                    continue

                if raw_code is None or raw_types is None:
                    raise CommandError(
                        f"Incomplete data at Excel row {row_number}"
                    )

                service_code = str(raw_code).strip().upper()
                cafe_types = [
                    item.strip().upper()
                    for item in str(raw_types).split(",")
                ]

                for cafe_type in cafe_types:
                    if cafe_type not in CAFE_CODE_MAP:
                        raise CommandError(
                            f"Unknown cafe type '{cafe_type}' "
                            f"at row {row_number}"
                        )

                    mappings.add(
                        (service_code, CAFE_CODE_MAP[cafe_type])
                    )

                service_codes.add(service_code)

            if len(service_codes) != 28:
                raise CommandError(
                    f"Expected 28 services, found {len(service_codes)}."
                )

            services = {
                item.code: item
                for item in GovernmentService.objects.filter(
                    code__in=service_codes
                )
            }

            activities = {
                item.code: item
                for item in Activity.objects.filter(
                    code__in=CAFE_CODE_MAP.values()
                )
            }

            missing_services = service_codes - set(services)
            missing_activities = set(CAFE_CODE_MAP.values()) - set(activities)

            if missing_services or missing_activities:
                raise CommandError(
                    f"Missing services: {sorted(missing_services)}; "
                    f"Missing activities: {sorted(missing_activities)}"
                )

            created_count = 0
            existing_count = 0

            with transaction.atomic():
                for service_code, activity_code in sorted(mappings):
                    _, created = ServiceActivity.objects.get_or_create(
                        service=services[service_code],
                        activity=activities[activity_code],
                    )

                    if created:
                        created_count += 1
                    else:
                        existing_count += 1

            self.stdout.write(
                self.style.SUCCESS(
                    f"Imported successfully! "
                    f"New relationships: {created_count}, "
                    f"Already existing: {existing_count}, "
                    f"Total in Excel: {len(mappings)}"
                )
            )

        finally:
            workbook.close()
