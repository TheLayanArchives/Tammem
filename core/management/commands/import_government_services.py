
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from core.models import Authority, GovernmentService


class Command(BaseCommand):
    help = "Import the 28 government services from the reviewed Excel file."

    def add_arguments(self, parser):
        parser.add_argument(
            "--apply",
            action="store_true",
            help="Save the validated services to the database.",
        )

    def handle(self, *args, **options):
        excel_path = (
            Path(settings.BASE_DIR)
            / "core"
            / "data"
            / "TAMMEM_28_Services_Reviewed_Decisions.xlsx"
        )

        if not excel_path.is_file():
            raise CommandError(f"Excel file not found: {excel_path}")

        workbook = load_workbook(
            excel_path, read_only=True, data_only=True
        )

        try:
            services = self.read_sheet(workbook, "Services")
            decisions = self.read_sheet(workbook, "Reviewed_Decisions")
        finally:
            workbook.close()

        service_codes = set(services)
        decision_codes = set(decisions)

        if service_codes != decision_codes:
            raise CommandError(
                "Service codes do not match between the two sheets. "
                f"Missing decisions: {sorted(service_codes - decision_codes)}. "
                f"Missing services: {sorted(decision_codes - service_codes)}."
            )

        if len(services) != 28:
            raise CommandError(
                f"Expected 28 services, found {len(services)}."
            )

        valid_visibility = {"ALWAYS", "CONDITIONAL"}
        valid_mandatory = {
            "REQUIRED", "OPTIONAL", "CONDITIONAL", "UNDECIDED"
        }

        prepared = []

        for code, row in sorted(services.items()):
            decision = decisions[code]

            authority_code = str(
                row.get("Authority_Code") or ""
            ).strip()

            visibility = str(
                decision.get("Visibility") or ""
            ).strip().upper()

            mandatory = str(
                decision.get("Mandatory_Status") or ""
            ).strip().upper()

            name_en = str(row.get("Name_EN") or "").strip()
            name_ar = str(row.get("Name_AR") or "").strip()

            if not name_en or not name_ar:
                raise CommandError(f"{code}: Missing service name.")

            if visibility not in valid_visibility:
                raise CommandError(
                    f"{code}: Invalid visibility: {visibility}"
                )

            if mandatory not in valid_mandatory:
                raise CommandError(
                    f"{code}: Invalid mandatory status: {mandatory}"
                )

            try:
                authority = Authority.objects.get(
                    code=authority_code
                )
            except Authority.DoesNotExist:
                raise CommandError(
                    f"{code}: Authority not found: {authority_code}"
                )

            prepared.append({
                "code": code,
                "authority": authority,
                "name_en": name_en,
                "name_ar": name_ar,
                "description_en": str(
                    row.get("Description_EN_Draft") or ""
                ).strip(),
                "official_url": str(
                    row.get("Official_URL") or ""
                ).strip(),
                "applicability": visibility,
                "mandatory_status": mandatory,
            })

        self.stdout.write(
            self.style.SUCCESS(
                f"Validation successful: {len(prepared)} services."
            )
        )

        for service in prepared:
            self.stdout.write(
                f"{service['code']} | "
                f"{service['applicability']} | "
                f"{service['mandatory_status']}"
            )

        if not options["apply"]:
            self.stdout.write(
                self.style.WARNING(
                    "PREVIEW ONLY: No database changes were made."
                )
            )
            self.stdout.write(
                "Use --apply to save the services."
            )
            return

        with transaction.atomic():
            for service in prepared:
                code = service["code"]

                defaults = {
                    key: value
                    for key, value in service.items()
                    if key != "code"
                }

                GovernmentService.objects.update_or_create(
                    code=code,
                    defaults=defaults,
                )

        self.stdout.write(
            self.style.SUCCESS(
                f"Successfully imported {len(prepared)} services."
            )
        )

    def read_sheet(self, workbook, sheet_name):
        if sheet_name not in workbook.sheetnames:
            raise CommandError(
                f"Missing Excel sheet: {sheet_name}"
            )

        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)

        headers = [
            str(value).strip() if value is not None else ""
            for value in next(rows)
        ]

        result = {}

        for row in rows:
            record = dict(zip(headers, row))
            code = str(
                record.get("Service_Code") or ""
            ).strip()

            if not code:
                continue

            if code in result:
                raise CommandError(
                    f"Duplicate service code in {sheet_name}: {code}"
                )

            result[code] = record

        return result
